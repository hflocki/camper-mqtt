"""Manufacturing files for PCBWay from the finished board.

usage: python3 make_outputs.py board.kicad_pcb outdir
Writes: gerber zip (copper, mask, paste, silk, outline, drill), BOM (xlsx + csv), placement file (csv).
"""
import collections
import csv
import os
import subprocess
import sys
import zipfile

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

from parts import PARTS

NAME = 'camper-heater-lin'
board, out = sys.argv[1], sys.argv[2]
os.makedirs(out, exist_ok=True)
tmp = os.path.join(out, '_gerber')
os.makedirs(tmp, exist_ok=True)

# ---- Gerber + drill --------------------------------------------------------------------------------------
layers = 'F.Cu,B.Cu,F.Paste,B.Paste,F.Silkscreen,B.Silkscreen,F.Mask,B.Mask,Edge.Cuts'
subprocess.run(['kicad-cli', 'pcb', 'export', 'gerbers', '-l', layers, '--subtract-soldermask', '--use-drill-file-origin', '-o', tmp + '/', board],
               check=True, capture_output=True)
subprocess.run(['kicad-cli', 'pcb', 'export', 'drill', '--format', 'excellon', '--excellon-separate-th',
                '--drill-origin', 'plot', '--generate-map', '--map-format', 'gerberx2', '-o', tmp + '/', board], check=True, capture_output=True)
zpath = os.path.join(out, NAME + '-gerber.zip')
with zipfile.ZipFile(zpath, 'w', zipfile.ZIP_DEFLATED) as z:
    for f in sorted(os.listdir(tmp)):
        z.write(os.path.join(tmp, f), f)
print(zpath, sorted(os.listdir(tmp)))

# ---- Placement (centroid) ----------------------------------------------------------------------------------
NOT_PLACED = lambda r: r.startswith(('H', 'TP', 'FID', 'JP'))
raw = os.path.join(tmp, 'pos.csv')
subprocess.run(['kicad-cli', 'pcb', 'export', 'pos', '--format', 'csv', '--units', 'mm', '--side', 'both',
                '--use-drill-file-origin', '-o', raw, board], check=True, capture_output=True)
rows = list(csv.DictReader(open(raw)))
cpl = os.path.join(out, NAME + '-cpl.csv')
with open(cpl, 'w', newline='') as f:
    w = csv.writer(f)
    w.writerow(['Designator', 'Mid X', 'Mid Y', 'Layer', 'Rotation'])
    for r in rows:
        if NOT_PLACED(r['Ref']):
            continue
        w.writerow([r['Ref'], r['PosX'] + 'mm', r['PosY'] + 'mm', 'Top' if r['Side'] == 'top' else 'Bottom', r['Rot']])
print(cpl, len(rows))

# ---- BOM -------------------------------------------------------------------------------------------------
THT = {'J1', 'J2', 'J3', 'F1', 'PS1'}
groups = collections.OrderedDict()
for ref, (lib, fp, val, mpn, mfr, desc, _) in PARTS.items():
    if NOT_PLACED(ref):
        continue
    key = (mpn, val, fp)
    if ref[0] == 'R' and ref[1].isdigit():
        desc = 'Widerstand 1 %, 0603'
    elif ref[0] == 'C' and ref[1].isdigit():
        desc = 'Keramikkondensator'
    groups.setdefault(key, {'refs': [], 'mfr': mfr, 'desc': desc, 'lib': lib})['refs'].append(ref)


def refsort(r):
    import re
    m = re.match(r'([A-Z]+)(\d+)', r)
    return (m.group(1), int(m.group(2)))


header = ['Item #', 'Designator', 'Qty', 'Manufacturer', 'Mfg Part #', 'Description / Value', 'Package / Footprint', 'Type', 'Hinweis']
lines = []
for i, ((mpn, val, fp), g) in enumerate(sorted(groups.items(), key=lambda kv: refsort(sorted(kv[1]['refs'], key=refsort)[0])), 1):
    refs = sorted(g['refs'], key=refsort)
    typ = 'THT' if refs[0] in THT else 'SMD'
    note = ''
    if refs[0] == 'F1':
        note = 'Nur der Halter wird bestückt. Sicherung 1 A (Mini-Flachsicherung ATM, z. B. Littelfuse 0297001.WXNV) lose beilegen oder selbst einstecken.'
    if refs[0] == 'U1':
        note = 'Antenne zeigt zum Platinenrand, unter der Antenne kein Kupfer.'
    if refs[0] == 'U2':
        note = 'Gleichwertig: TJA1021T/20/C,118 oder TJA1021T/10 (LIN 2.x Slave).'
    lines.append([i, ','.join(refs), len(refs), g['mfr'], mpn, f'{val} – {g["desc"]}', fp, typ, note])

bom_csv = os.path.join(out, NAME + '-bom.csv')
with open(bom_csv, 'w', newline='') as f:
    w = csv.writer(f)
    w.writerow(header)
    w.writerows(lines)
wb = openpyxl.Workbook()
ws = wb.active
ws.title = 'BOM'
ws.append(header)
for c in ws[1]:
    c.font = Font(bold=True)
    c.fill = PatternFill('solid', fgColor='DDDDDD')
for l in lines:
    ws.append(l)
for col, width in zip('ABCDEFGHI', (7, 22, 6, 20, 26, 52, 44, 7, 60)):
    ws.column_dimensions[col].width = width
for row in ws.iter_rows(min_row=2):
    for c in row:
        c.alignment = Alignment(wrap_text=True, vertical='top')
bom = os.path.join(out, NAME + '-bom.xlsx')
wb.save(bom)
print(bom, len(lines), 'lines')

for f in os.listdir(tmp):
    os.remove(os.path.join(tmp, f))
os.rmdir(tmp)
